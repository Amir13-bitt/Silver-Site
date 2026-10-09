#!/usr/bin/env python3
"""Convert the store's product export (.xlsx) into js/products.js.

Usage:  python3 scripts/import_products.py path/to/products.xlsx

Each product has two dimensions: a category (from the "دسته‌بندی‌ها" column)
and a collection (from the "مجموعه" column). When the collection column is
empty, the collection is detected from the description, and COLLECTION_OVERRIDES
below always wins, so products can be assigned by hand.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

# handle -> collection key. Fill in products whose collection isn't in the export.
COLLECTION_OVERRIDES = {
    # "horse-fancy-ring": "tasyan",
}

# English names/descriptions, keyed by handle. Persian comes from the export.
EN = {
    "lavan-earrings": ("Lavan Earrings", "Inspired by the nature of southern Iran."),
    "gulf-earrings": ("Gulf Earrings", "Inspired by the nature of southern Iran."),
    "tt-earrings": ("Titi Earrings", "The experiences we carry with us through life. Made by hand, every time."),
    "ahil-and-pearl-earrings": ("Ahil & Pearl Earrings", "Traces of the moments that stay with us. Made by hand, every time. Freshwater pearl."),
    "ircaf-silver-douran": ("Douran Silver Ear Cuff", "3 grams."),
    "still-necklace": ("Still Necklace", "A silver necklace from the Flower Power collection, shaped after urban flowers: irregular, alive, and not trying to be perfect. A sign of growth and carrying on, despite everything. Available on a 40 or 50 cm silver chain."),
    "darakoh-silver-ring": ("Dar-e Kuh Silver Ring", "The “Dar-e Kuh” ring from the Chaos collection, inspired by the nature of southern Iran."),
    "silver-smooth-earrings": ("Ravan Silver Earrings", "From the Flower Power collection: a flower-like form with fluid lines that suggest free movement. Its hollow shape and soft edges let light pass through."),
    "invisible-silver-earrings": ("Napeyda Silver Earrings", "From the Flower Power collection: a trace that may be invisible, but is never without effect."),
    "stud-earring": ("Yal Earrings", "Silver and freshwater pearl."),
    "double-sided-earrings": ("Face to Face Earrings", "Silver."),
    "chinese-cloud-chest-stud": ("Chinese Cloud Brooch", "Silver."),
    "once-upon-a-time-earring": ("Once Upon a Time Earrings", "Silver."),
    "carousel-necklace": ("Carousel Necklace", "Silver."),
    "tuscan-earrings": ("Tosan Earrings", "Silver."),
    "horse-fancy-ring": ("Dream of a Horse Ring", "Silver."),
    "mint-ring": ("Navand Ring", "Silver."),
    "horse-necklace": ("A Horse Necklace", "Silver."),
    "chocolate-cake": ("Galloping Choker", "Silver and brass cord."),
}

CATEGORIES = {
    "Earrings & Ear Cuffs": "earrings",
    "Necklaces & Pendants": "necklaces",
    "Rings": "rings",
    "Brooches & Pins": "brooches",
}

COLLECTION_PATTERNS = [
    ("flower-power", re.compile(r"flower\s*power", re.I)),
    ("chaos", re.compile(r"chaos|کائوس|کیاس", re.I)),
    ("tasyan", re.compile(r"tasyan|تسیان", re.I)),
]


def collection_for(handle, explicit, description):
    if handle in COLLECTION_OVERRIDES:
        return COLLECTION_OVERRIDES[handle]
    for key, pattern in COLLECTION_PATTERNS:
        if explicit and pattern.search(str(explicit)):
            return key
    for key, pattern in COLLECTION_PATTERNS:
        if pattern.search(description or ""):
            return key
    return None


def main(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    header = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(header)}
    image_cols = [i for i, name in enumerate(header) if name and str(name).startswith("تصویر ")]

    products = []
    for row in ws.iter_rows(min_row=2):
        v = [c.value for c in row]
        if v[col["وضعیت"]] != "منتشر شده":
            continue  # drafts are hidden on the store too
        handle = v[col["نامک (Handle)"]]
        description = (v[col["توضیحات"]] or "").strip()
        category_raw = (v[col["دسته‌بندی‌ها"]] or "").split("›")[-1].strip()
        price = int(re.sub(r"[^\d]", "", str(v[col["قیمت"]]) or "0") or 0)
        stock = v[col["موجودی"]]
        images = [
            row[i].hyperlink.target
            for i in image_cols
            if row[i].hyperlink and not row[i].hyperlink.target.startswith("http://localhost")
        ]
        name_en, desc_en = EN.get(handle, (handle.replace("-", " ").title(), ""))
        products.append({
            "handle": handle,
            "name": {"fa": v[col["محصول"]].strip(), "en": name_en},
            "desc": {"fa": description, "en": desc_en},
            "category": CATEGORIES.get(category_raw, "other"),
            "collection": collection_for(handle, v[col["مجموعه"]], description),
            "price": price,  # IRR
            "stock": stock if isinstance(stock, int) else None,
            "images": images,
        })

    out = Path(__file__).resolve().parent.parent / "js" / "products.js"
    out.write_text(
        "/* Generated by scripts/import_products.py from the store export. Do not edit by hand. */\n"
        "window.PRODUCTS = " + json.dumps(products, ensure_ascii=False, indent=2) + ";\n",
        encoding="utf-8",
    )
    found = sum(1 for p in products if p["collection"])
    print(f"Wrote {len(products)} products to {out} ({found} with a collection).")


if __name__ == "__main__":
    main(sys.argv[1])
